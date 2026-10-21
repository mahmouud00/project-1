import pandas as pd

pd.set_option('display.max_columns', None)
pd.set_option('display.max_rows', None)
import numpy as np
import openpyxl as xl
from openpyxl.chart import BarChart, Reference

df = pd.read_excel("C:/Users/Mahmoud/Desktop/مراجعة_عملاء_صيانة.xlsx", header=1)
df = df.dropna(how='all')
print(df.head(1))

parts = df["الاسم||رقم الموبايل"].str.split(r"\|+", expand=True, regex=True)
parts[0] = parts[0].str.strip()
parts[1] = parts[1].str.strip()
df["الرقم"] = np.where(parts[0].str.isdigit(), parts[0], parts[1])
df["الاسم"] = np.where(parts[1].str.isdigit(), parts[0], parts[1])
rush_day = df._series["التاريخ"]
device = df._series["الجهاز"]
most_devices = df.groupby(["التاريخ", "الجهاز"]).size()
follow_up = df[(df["الرد"].isna()) | (df["call"].isna())]
wb = xl.load_workbook("C:/Users/Mahmoud/Desktop/مراجعة_عملاء_صيانة.xlsx")
sheet = wb["Sheet1"]
ref = Reference(sheet, 3, 9, 3, sheet.max_row)
chart = BarChart()
chart.add_data(ref)
sheet.add_chart(chart, "G9")
print(follow_up)
print(most_devices)
print(rush_day.value_counts())
print(device.value_counts())
print(df["الاسم"], df["الرقم"])
